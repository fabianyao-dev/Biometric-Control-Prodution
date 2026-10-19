"""
test_export_auto.py - Exportacion automatica (snapshot por hora + cierre diario).

Dos niveles:

  - La DECISION (`ruta_export`, `decidir`, `decidir_hoy`) es pura: se prueba
    con datos fijos, sin reloj real, sin BD y sin escribir archivos.
  - El programador con providers inyectados (`ProgramadorExport`) se prueba
    con un reloj y un exportar falsos, y un caso end-to-end contra la BD
    temporal con el `exportar_dia` real.
"""

import os
from datetime import datetime

import pytest

from src.reportes.export_auto import (
    DIAS_CATCHUP_MAX,
    EstadoExport,
    ProgramadorExport,
    TrabajoExport,
    decidir,
    decidir_hoy,
    ruta_export,
)

# ---------------------------------------------------------------------------
# Helpers puros
# ---------------------------------------------------------------------------


def _dt(texto):
    return datetime.strptime(texto, "%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------------------
# ruta_export
# ---------------------------------------------------------------------------


def test_ruta_export_usa_una_carpeta_por_dia():
    base = r"C:\datos\export"
    final = ruta_export(base, "2026-10-17", parcial=False)
    parcial = ruta_export(base, "2026-10-17", parcial=True)

    carpeta = os.path.join(base, "2026-10-17")
    assert final == os.path.join(carpeta, "Sesiones_2026-10-17.xlsx")
    assert parcial == os.path.join(carpeta, "Sesiones_2026-10-17_parcial.xlsx")
    # Dias distintos -> carpetas distintas.
    otro = ruta_export(base, "2026-10-18", parcial=False)
    assert os.path.dirname(otro) != os.path.dirname(final)


# ---------------------------------------------------------------------------
# decidir (pura)
# ---------------------------------------------------------------------------


def test_decidir_sin_actividad_no_hace_nada():
    trabajos = decidir(
        _dt("2026-10-17 09:00:00"), EstadoExport(),
        cada_hora=True, hoy="2026-10-17",
        hoy_tiene_sesiones=False, limite="2026-10-16", dias_pendientes=[],
    )
    assert trabajos == []


def test_decidir_snapshot_horario_una_vez_por_hora():
    ahora = _dt("2026-10-17 09:00:00")
    trabajos = decidir(
        ahora, EstadoExport(), cada_hora=True,
        hoy="2026-10-17", hoy_tiene_sesiones=True, limite="2026-10-16",
        dias_pendientes=[],
    )
    assert [(t.tipo, t.fecha, t.parcial) for t in trabajos] == [
        ("hora", "2026-10-17", True)]

    # Ya hecho a esta hora -> no se repite.
    trabajos = decidir(
        ahora, EstadoExport(ultima_hora="2026-10-17 09"), cada_hora=True,
        hoy="2026-10-17", hoy_tiene_sesiones=True,
        limite="2026-10-16", dias_pendientes=[],
    )
    assert trabajos == []


def test_decidir_cierra_un_dia_cerrado_pero_no_uno_abierto():
    ahora = _dt("2026-10-17 01:00:00")
    # Dia con sesiones y ya cerrado -> se exporta.
    trabajos = decidir(
        ahora, EstadoExport(), cada_hora=False,
        hoy="2026-10-17", hoy_tiene_sesiones=False, limite="2026-10-16",
        dias_pendientes=[("2026-10-15", True, True), ("2026-10-16", True, True)],
    )
    assert [(t.tipo, t.fecha) for t in trabajos] == [
        ("catchup", "2026-10-15"), ("dia", "2026-10-16")]

    # Dia con sesion ABIERTA -> no se exporta (se reintenta luego).
    trabajos = decidir(
        ahora, EstadoExport(), cada_hora=False,
        hoy="2026-10-17", hoy_tiene_sesiones=False, limite="2026-10-16",
        dias_pendientes=[("2026-10-16", True, False)],
    )
    assert trabajos == []

    # Dia sin sesiones -> no hay archivo que hacer.
    trabajos = decidir(
        ahora, EstadoExport(), cada_hora=False,
        hoy="2026-10-17", hoy_tiene_sesiones=False, limite="2026-10-16",
        dias_pendientes=[("2026-10-16", False, True)],
    )
    assert trabajos == []


# ---------------------------------------------------------------------------
# decidir_hoy (pura)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("tiene,cerrado,esperado", [
    (True, True, ("hoy", "2026-10-17")),   # ultima sesion del dia cerrada
    (True, False, None),                   # queda una sesion abierta
    (False, True, None),                   # dia sin sesiones todavia
])
def test_decidir_hoy(tiene, cerrado, esperado):
    trabajo = decidir_hoy("2026-10-17", tiene, cerrado)
    if esperado is None:
        assert trabajo is None
    else:
        assert (trabajo.tipo, trabajo.fecha) == esperado
        assert trabajo.parcial is False  # es el cierre, no el snapshot


# ---------------------------------------------------------------------------
# ProgramadorExport con providers inyectados
# ---------------------------------------------------------------------------


class _FakeExport:
    """Registra (fecha, ruta) y crea el archivo para que el replace funcione."""

    def __init__(self, ok=True):
        self.llamadas = []
        self.ok = ok

    def __call__(self, fecha, ruta):
        self.llamadas.append((fecha, ruta))
        if self.ok:
            with open(ruta, "wb") as f:
                f.write(b"xlsx")
        return self.ok, f"ok {fecha}", {}


def _programador(tmp_path, *, ahora, sesiones=(), cerrados=True, cada_hora=True,
                 ultimo_dia="", exportar=None, persistir=None):
    """Programador con reloj/providers fijos.

    `sesiones` es el conjunto de fechas que SI tienen sesiones; `cerrados`
    indica si esos dias estan cerrados (sin sesion activa).
    """
    fechas = set(sesiones)
    return ProgramadorExport(
        dir_auto=str(tmp_path),
        cada_hora=cada_hora,
        reloj=lambda: ahora,
        tiene_sesiones=lambda f: f in fechas,
        esta_cerrado=lambda f: cerrados,
        exportar=exportar or _FakeExport(),
        leer_estado=lambda: ultimo_dia,
        persistir=persistir or (lambda f: None),
    )


def test_pendientes_cierra_ayer_sin_esperar_ninguna_hora(tmp_path):
    """El fin de turno varia, asi que ayer se cierra apenas queda cerrado."""
    for hora in ("00:01:00", "09:30:00", "22:00:00"):
        prog = _programador(
            tmp_path, ahora=_dt(f"2026-10-17 {hora}"),
            sesiones={"2026-10-15", "2026-10-16"}, cada_hora=False,
            ultimo_dia="2026-10-01",
        )
        assert [t.fecha for t in prog.pendientes()] == [
            "2026-10-15", "2026-10-16"], f"fallo a las {hora}"

    # Ayer con una sesion todavia abierta -> no se cierra (se reintenta).
    abierto = _programador(
        tmp_path, ahora=_dt("2026-10-17 22:00:00"),
        sesiones={"2026-10-16"}, cerrados=False, cada_hora=False,
        ultimo_dia="2026-10-15",
    )
    assert abierto.pendientes() == []


def test_pendientes_incluir_hoy_solo_tras_cerrar_la_ultima_sesion(tmp_path):
    """El tick no toca hoy; el cierre de sesion/app si lo pide."""
    cerrado = _programador(
        tmp_path, ahora=_dt("2026-10-17 17:00:00"),
        sesiones={"2026-10-17"}, cada_hora=False, ultimo_dia="2026-10-16",
    )
    # Por defecto (tick) no sale nada: hoy no se cierra por reloj.
    assert cerrado.pendientes() == []
    # Con el aviso de cierre de sesion, sale el cierre de hoy.
    trabajos = cerrado.pendientes(incluir_hoy=True)
    assert [(t.tipo, t.fecha, t.parcial) for t in trabajos] == [
        ("hoy", "2026-10-17", False)]

    # Con una sesion de hoy todavia abierta -> no hay cierre que hacer.
    abierto = _programador(
        tmp_path, ahora=_dt("2026-10-17 12:00:00"),
        sesiones={"2026-10-17"}, cerrados=False, cada_hora=False,
        ultimo_dia="2026-10-16",
    )
    assert abierto.pendientes(incluir_hoy=True) == []

    # Un dia sin sesiones no genera cierre aunque se pida.
    vacio = _programador(
        tmp_path, ahora=_dt("2026-10-17 12:00:00"),
        sesiones=set(), cada_hora=False, ultimo_dia="2026-10-16",
    )
    assert vacio.pendientes(incluir_hoy=True) == []


def test_pendientes_respeta_el_ultimo_dia_exportado(tmp_path):
    prog = _programador(
        tmp_path, ahora=_dt("2026-10-17 09:00:00"),
        sesiones={"2026-10-14", "2026-10-15", "2026-10-16"}, cada_hora=False,
        ultimo_dia="2026-10-15",
    )
    # Ya se exporto hasta el 15: solo queda el 16.
    assert [t.fecha for t in prog.pendientes()] == ["2026-10-16"]


def test_pendientes_primer_arranque_no_reexporta_todo(tmp_path):
    """Sin historial solo se mira el dia limite, no los 14 anteriores."""
    prog = _programador(
        tmp_path, ahora=_dt("2026-10-17 09:00:00"),
        sesiones={f"2026-10-{d:02d}" for d in range(1, 17)}, cada_hora=False,
        ultimo_dia="",
    )
    assert [t.fecha for t in prog.pendientes()] == ["2026-10-16"]


def test_pendientes_acota_el_catchup(tmp_path):
    prog = _programador(
        tmp_path, ahora=_dt("2026-10-17 09:00:00"),
        sesiones={f"2026-10-{d:02d}" for d in range(1, 17)}, cada_hora=False,
        ultimo_dia="2026-09-01",
    )
    fechas = [t.fecha for t in prog.pendientes()]
    assert len(fechas) == DIAS_CATCHUP_MAX
    assert fechas[-1] == "2026-10-16"


def test_apagado_no_produce_trabajos(tmp_path):
    prog = _programador(
        tmp_path, ahora=_dt("2026-10-17 09:00:00"),
        sesiones={"2026-10-17"}, cada_hora=True,
    )
    prog.activo = False
    assert prog.pendientes() == []


# ---------------------------------------------------------------------------
# ejecutar / confirmar
# ---------------------------------------------------------------------------


def test_ejecutar_escribe_el_archivo_final(tmp_path):
    fake = _FakeExport()
    prog = _programador(tmp_path, ahora=_dt("2026-10-17 01:00:00"),
                        exportar=fake)
    ok, msg = prog.ejecutar(TrabajoExport("dia", "2026-10-16", False))

    assert ok, msg
    destino = tmp_path / "2026-10-16" / "Sesiones_2026-10-16.xlsx"
    assert destino.exists()
    # El export escribio a un temporal y se reemplazo: no queda basura.
    assert not (tmp_path / "2026-10-16"
                / "Sesiones_2026-10-16.xlsx.tmp.xlsx").exists()
    assert fake.llamadas[0][0] == "2026-10-16"


def test_ejecutar_no_escribe_si_el_export_falla(tmp_path):
    fake = _FakeExport(ok=False)
    prog = _programador(tmp_path, ahora=_dt("2026-10-17 01:00:00"),
                        exportar=fake)
    ok, _ = prog.ejecutar(TrabajoExport("dia", "2026-10-16", False))

    assert not ok
    # No queda ningun archivo del dia (la carpeta puede quedar vacia).
    assert not list(tmp_path.glob("**/Sesiones_2026-10-16*"))


def test_ejecutar_destino_bloqueado_guarda_copia_con_hora(tmp_path, monkeypatch):
    """Si Excel tiene abierto el destino, no se pierde el reporte."""
    prog = _programador(tmp_path, ahora=_dt("2026-10-17 10:30:00"))
    destino = (tmp_path / "2026-10-17"
               / "Sesiones_2026-10-17_parcial.xlsx")

    import src.reportes.export_auto as modulo

    real_replace = modulo.os.replace

    def replace_bloqueado(origen, fin):
        if str(fin) == str(destino):
            raise PermissionError("abierto en Excel")
        return real_replace(origen, fin)

    monkeypatch.setattr(modulo.os, "replace", replace_bloqueado)

    ok, msg = prog.ejecutar(TrabajoExport("hora", "2026-10-17", True))

    assert ok, msg
    assert "abierto" in msg
    assert (tmp_path / "2026-10-17"
            / "Sesiones_2026-10-17_parcial_1030.xlsx").exists()


def test_confirmar_avanza_el_estado_y_persiste(tmp_path):
    persistidos = []
    prog = _programador(tmp_path, ahora=_dt("2026-10-17 09:20:00"),
                        persistir=persistidos.append)

    prog.confirmar(TrabajoExport("hora", "2026-10-17", True))
    assert prog.estado.ultima_hora == "2026-10-17 09"
    assert persistidos == []  # el snapshot no cambia el ultimo dia

    prog.confirmar(TrabajoExport("dia", "2026-10-16", False))
    assert prog.estado.ultimo_dia == "2026-10-16"
    assert persistidos == ["2026-10-16"]


def test_confirmar_hoy_no_sella_el_dia(tmp_path):
    """El cierre del dia en curso NO avanza `ultimo_dia`: hoy sigue abierto.

    Asi, si el cierre de anoche fallo o hubo un receso, el dia se reexporta
    al dia siguiente, ya terminado, y recien entonces se sella.
    """
    persistidos = []
    prog = _programador(tmp_path, ahora=_dt("2026-10-17 17:00:00"),
                        ultimo_dia="2026-10-16", persistir=persistidos.append)

    prog.confirmar(TrabajoExport("hoy", "2026-10-17", False))

    assert prog.estado.ultimo_dia == "2026-10-16"
    assert persistidos == []


def test_confirmar_no_retrocede_el_ultimo_dia(tmp_path):
    prog = _programador(tmp_path, ahora=_dt("2026-10-17 09:00:00"),
                        ultimo_dia="2026-10-16")
    prog.confirmar(TrabajoExport("catchup", "2026-10-10", False))
    assert prog.estado.ultimo_dia == "2026-10-16"


# ---------------------------------------------------------------------------
# End-to-end contra la BD temporal
# ---------------------------------------------------------------------------


def _crear_operador(database, nombre="Ana Reporte"):
    for rol in database.listar_roles():
        if rol["nombre"] == "operador":
            rol_id = rol["id"]
            break
    else:
        raise AssertionError("falta el rol operador")
    ok, resultado = database.guardar_operador(
        nombre, bytes(range(16)) * 3 + b"\x00\x01", rol_id)
    assert ok, f"no se pudo crear el operador: {resultado}"
    return resultado


def test_snapshot_horario_end_to_end(db, tmp_path):
    """Con la BD real: el snapshot por hora genera el .xlsx del dia en curso."""
    from openpyxl import load_workbook

    import src.database as database

    operador_id = _crear_operador(database)
    sesion_id = database.abrir_sesion(operador_id)
    database.cerrar_sesion(sesion_id, 7)

    prog = ProgramadorExport.desde_config()
    prog.dir_auto = str(tmp_path)
    prog.cada_hora = True
    prog.estado = EstadoExport()  # sin historial

    trabajos = prog.pendientes()
    assert trabajos, "deberia tocar el snapshot por hora"
    trabajo = trabajos[0]
    assert trabajo.tipo == "hora" and trabajo.parcial

    ok, msg = prog.ejecutar(trabajo)
    assert ok, msg

    ruta = (tmp_path / trabajo.fecha
            / f"Sesiones_{trabajo.fecha}_parcial.xlsx")
    assert ruta.exists()
    wb = load_workbook(ruta)
    try:
        assert wb["Sesiones"].max_row == 2
    finally:
        wb.close()


def test_cierre_de_hoy_end_to_end(db, tmp_path):
    """Con la BD real: el cierre del dia en curso genera el .xlsx sin sufijo.

    Es el archivo que sale al cerrar la ultima sesion del dia (o al cerrar la
    app), antes de que se apague la PC; el tick por si solo no lo pide.
    """
    from openpyxl import load_workbook

    import src.database as database

    operador_id = _crear_operador(database)
    sesion_id = database.abrir_sesion(operador_id)
    database.cerrar_sesion(sesion_id, 9)

    prog = ProgramadorExport.desde_config()
    prog.dir_auto = str(tmp_path)
    prog.cada_hora = False   # aislar el cierre de hoy del snapshot por hora
    prog.estado = EstadoExport(ultimo_dia="2026-01-01")  # historial viejo

    # El tick no incluye hoy; el cierre de sesion/app si.
    assert prog.pendientes() == []
    trabajos = prog.pendientes(incluir_hoy=True)
    assert [(t.tipo, t.parcial) for t in trabajos] == [("hoy", False)]
    trabajo = trabajos[0]

    ok, msg = prog.ejecutar(trabajo)
    assert ok, msg

    ruta = tmp_path / trabajo.fecha / f"Sesiones_{trabajo.fecha}.xlsx"
    assert ruta.exists()
    assert not list(tmp_path.glob("**/*_parcial.xlsx"))
    wb = load_workbook(ruta)
    try:
        assert wb["Sesiones"].max_row == 2
    finally:
        wb.close()

    # Un cierre de hoy exitoso NO sella el dia: sigue abierto para el catch-up.
    prog.confirmar(trabajo)
    assert prog.estado.ultimo_dia == "2026-01-01"
